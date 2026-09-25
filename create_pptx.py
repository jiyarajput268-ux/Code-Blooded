import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]
    
    # Color Palette
    BURGUNDY = RGBColor(0x58, 0x1C, 0x25)
    CREAM = RGBColor(0xFD, 0xFB, 0xF7)
    GOLD = RGBColor(0xF4, 0xE8, 0xC1)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    DARK_CARD = RGBColor(0x4A, 0x15, 0x21)
    ACCENT_GREEN = RGBColor(0x10, 0xB9, 0x81)
    ACCENT_ROSE = RGBColor(0xE1, 0x1D, 0x48)
    ACCENT_AMBER = RGBColor(0xF5, 0x9E, 0x0B)
    
    def apply_bg(slide, color=BURGUNDY):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, slide_num, title_text, category_text="AMIHACKS 1.0 • HACKATHON PRESENTATION (SECTION 6.2)"):
        # Header bar
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
        tf = txBox.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = category_text
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = GOLD
        
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(26)
        p1.font.bold = True
        p1.font.color.rgb = CREAM
        
        # Slide counter
        counterBox = slide.shapes.add_textbox(Inches(10.8), Inches(0.4), Inches(1.8), Inches(0.5))
        c_tf = counterBox.text_frame
        c_p = c_tf.paragraphs[0]
        c_p.text = f"Slide {slide_num} of 8"
        c_p.alignment = PP_ALIGN.RIGHT
        c_p.font.size = Pt(12)
        c_p.font.bold = True
        c_p.font.color.rgb = GOLD

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    apply_bg(slide1, BURGUNDY)
    
    card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    card1.fill.solid()
    card1.fill.fore_color.rgb = DARK_CARD
    card1.line.color.rgb = GOLD
    
    tb1 = slide1.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(10.9), Inches(5.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "🏆 AMIHACKS 1.0 HACKATHON PITCH"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = GOLD
    
    p = tf1.add_paragraph()
    p.text = "PlateRelay: \"Surplus-to-Shelter\""
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = CREAM
    
    p = tf1.add_paragraph()
    p.text = "Real-Time Food Rescue Routing Platform"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = GOLD
    
    p = tf1.add_paragraph()
    p.text = "\nEliminating urban food waste via real-time 3-stage perishability cascade, automated NGO matching engine, and dual-OTP chain-of-custody verification."
    p.font.size = Pt(15)
    p.font.color.rgb = WHITE
    
    # Feature Badges
    badges = [
        ("3-Stage Cascade Engine", "Commercial -> Share -> NGO"),
        ("78 / 100 Match Score", "Dynamic Weighted Algorithm"),
        ("Dual-OTP Security", "#4821 Pickup | #9374 Delivery"),
        ("Dual DB Architecture", "SQLite + Mongo Atlas Sync")
    ]
    for i, (b_title, b_desc) in enumerate(badges):
        b_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2 + i * 2.7), Inches(4.8), Inches(2.5), Inches(1.2))
        b_box.fill.solid()
        b_box.fill.fore_color.rgb = BURGUNDY
        b_box.line.color.rgb = GOLD
        
        b_tf = b_box.text_frame
        b_tf.word_wrap = True
        bp1 = b_tf.paragraphs[0]
        bp1.text = b_title
        bp1.font.size = Pt(11)
        bp1.font.bold = True
        bp1.font.color.rgb = GOLD
        
        bp2 = b_tf.add_paragraph()
        bp2.text = b_desc
        bp2.font.size = Pt(9)
        bp2.font.color.rgb = CREAM

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    apply_bg(slide2, BURGUNDY)
    add_header(slide2, 2, "The Urban Food Paradox: Waste vs. Hunger", "🔴 PROBLEM STATEMENT")
    
    problems = [
        ("🚮 67 Million Tons Wasted", "India loses ₹92,000 Crores worth of edible food annually. Commercial kitchens, hostels, weddings, and bakeries throw away edible surplus daily."),
        ("⏳ Perishability Clock Barrier", "Cooked meals spoil within 2–4 hours. Traditional phone-call networks to NGOs fail because food expires before volunteers can arrive."),
        ("🚫 Lack of Smart Logistics", "Zero real-time unified platforms exist to route surplus food dynamically between cost-conscious buyers, community members, and shelters.")
    ]
    for i, (title, desc) in enumerate(problems):
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.9), Inches(1.8), Inches(3.7), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = DARK_CARD
        card.line.color.rgb = ACCENT_ROSE
        
        tf = card.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = GOLD
        
        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(13)
        p2.font.color.rgb = WHITE

    # -------------------------------------------------------------
    # SLIDE 3: Proposed Solution
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    apply_bg(slide3, BURGUNDY)
    add_header(slide3, 3, "PlateRelay: Real-Time 3-Stage Cascade System", "💡 PROPOSED SOLUTION")
    
    stages = [
        ("STAGE 1: RESCUE DEAL (50-70% OFF) 🏷️", "Commercial resale window allowing food businesses to recover preparation costs while food is super fresh.", ACCENT_AMBER),
        ("STAGE 2: FREE SHARE FEED 🌱", "Zero-cost community sharing feed activated automatically as safe consumption time window decreases.", ACCENT_GREEN),
        ("STAGE 3: NGO BULK RESCUE 🤝", "Automated direct routing to verified shelters & langars for bulk quantities (≥ 15 portions) or near cutoff time.", ACCENT_ROSE)
    ]
    for i, (stitle, sdesc, color) in enumerate(stages):
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.9), Inches(1.8), Inches(3.7), Inches(4.2))
        card.fill.solid()
        card.fill.fore_color.rgb = DARK_CARD
        card.line.color.rgb = color
        
        tf = card.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = stitle
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = color
        
        p2 = tf.add_paragraph()
        p2.text = f"\n{sdesc}"
        p2.font.size = Pt(13)
        p2.font.color.rgb = WHITE

    # Bottom Safety Clock Note
    clock_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.8))
    clock_card.fill.solid()
    clock_card.fill.fore_color.rgb = DARK_CARD
    clock_card.line.color.rgb = GOLD
    c_tf = clock_card.text_frame
    c_p = c_tf.paragraphs[0]
    c_p.text = "⏱️ Safety Countdown Clock Ring: Displays live safe-for timer (e.g., 3h 45m) ensuring zero spoiled food distribution."
    c_p.font.size = Pt(12)
    c_p.font.bold = True
    c_p.font.color.rgb = GOLD

    # -------------------------------------------------------------
    # SLIDE 4: Key Features
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    apply_bg(slide4, BURGUNDY)
    add_header(slide4, 4, "Intelligent Modules & Value Proposition", "✨ KEY FEATURES & INNOVATION")
    
    features = [
        ("🧠 Automated Weighted NGO Matcher", "Calculates dynamic match scores (78/100) using Distance (40%) + Capacity Fit (25%) + Urgency (20%) + Preference (15%)."),
        ("📷 Smart Dish Photo Matcher", "Auto-detects dish titles (Momos, Paneer, Biryani, Chole Bhature, Gulab Jamun, Dal Makhani) with direct device photo upload option."),
        ("🔒 Dual-OTP Handoff Security", "Pickup OTP (#4821) and Delivery OTP (#9374) verify volunteer runner chain-of-custody to eliminate fraud."),
        ("🗺️ OpenStreetMap Integration", "Live interactive map with custom Leaflet markers for Temples, Langars, and Verified Community Shelters.")
    ]
    for i, (ftitle, fdesc) in enumerate(features):
        row = i // 2
        col = i % 2
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + col * 5.95), Inches(1.8 + row * 2.6), Inches(5.75), Inches(2.4))
        card.fill.solid()
        card.fill.fore_color.rgb = DARK_CARD
        card.line.color.rgb = GOLD
        
        tf = card.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = ftitle
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = GOLD
        
        p2 = tf.add_paragraph()
        p2.text = f"\n{fdesc}"
        p2.font.size = Pt(12)
        p2.font.color.rgb = WHITE

    # -------------------------------------------------------------
    # SLIDE 5: System Architecture
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    apply_bg(slide5, BURGUNDY)
    add_header(slide5, 5, "Architecture & Tech Stack Breakdown", "🏗️ SYSTEM ARCHITECTURE")
    
    arch_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(2.2))
    arch_card.fill.solid()
    arch_card.fill.fore_color.rgb = DARK_CARD
    arch_card.line.color.rgb = GOLD
    a_tf = arch_card.text_frame
    a_tf.word_wrap = True
    ap = a_tf.paragraphs[0]
    ap.text = "[ Donor Form Input (React) ] ──► [ Express REST API (Port 5001) ] ──► [ Perishability Engine ]\n                                                     │\n                                                     ▼\n                                      ┌──────────────────────────┐\n                                      │ Dual-Database Architecture│\n                                      │ • SQLite (database.db)   │\n                                      │ • Mongo Atlas Cloud Sync │\n                                      └──────────────────────────┘"
    ap.font.size = Pt(13)
    ap.font.bold = True
    ap.font.color.rgb = GOLD

    stacks = [
        ("Frontend Layer", "React.js, Vite, Tailwind CSS, Lucide Icons, Leaflet Maps"),
        ("Backend & Engine", "Node.js, Express API, Haversine Distance, Weighted Matcher"),
        ("Database & Storage", "SQLite (Zero-setup local) + MongoDB Atlas Cloud Cluster")
    ]
    for i, (stitle, sdesc) in enumerate(stacks):
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.9), Inches(4.3), Inches(3.7), Inches(2.5))
        card.fill.solid()
        card.fill.fore_color.rgb = DARK_CARD
        card.line.color.rgb = CREAM
        
        tf = card.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = stitle
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = GOLD
        
        p2 = tf.add_paragraph()
        p2.text = f"\n{sdesc}"
        p2.font.size = Pt(12)
        p2.font.color.rgb = WHITE

    # -------------------------------------------------------------
    # SLIDE 6: Demo Screenshots & Gallery
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    apply_bg(slide6, BURGUNDY)
    add_header(slide6, 6, "Live Platform Workflows & Exact Dish Assets", "📸 DEMO SCREENSHOTS & GALLERY")
    
    shots = [
        ("1. Dal Makhani Rescue", "c:/Users/sharm/OneDrive/Desktop/Code Blooded/client/public/dal_makhani.png", "Exact uploaded Dal Makhani photo with cream swirl & butter."),
        ("2. Chole Bhature Rescue", "c:/Users/sharm/OneDrive/Desktop/Code Blooded/client/public/chole_bhature.png", "Fluffy bhaturas & paneer chole matched automatically."),
        ("3. Gulab Jamun & Sweets", "c:/Users/sharm/OneDrive/Desktop/Code Blooded/client/public/gulab_jamun.png", "Authentic sweet syrup Gulab Jamun with instant donor handoff.")
    ]
    for i, (title, img_path, desc) in enumerate(shots):
        card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.9), Inches(1.8), Inches(3.7), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = DARK_CARD
        card.line.color.rgb = GOLD
        
        tf = card.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = GOLD
        
        if os.path.exists(img_path):
            slide6.shapes.add_picture(img_path, Inches(1.0 + i * 3.9), Inches(2.4), Inches(3.3), Inches(2.5))
            
        p2 = tf.add_paragraph()
        p2.text = f"\n\n\n\n\n\n\n\n{desc}"
        p2.font.size = Pt(11)
        p2.font.color.rgb = WHITE

    # -------------------------------------------------------------
    # SLIDE 7: Technical Implementation & Impact
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    apply_bg(slide7, BURGUNDY)
    add_header(slide7, 7, "Simulated Clock & Environmental Analytics", "📊 TECHNICAL IMPLEMENTATION & IMPACT")
    
    impacts = [
        ("⏱️ Demo Clock Acceleration Simulator", "Built-in simulated time controller allowing hackathon judges to fast-forward hours in real-time and watch listings transition dynamically across Stage 1 → Stage 2 → Stage 3."),
        ("🌱 Environmental CO₂e Impact Formula", "CO₂e Saved (kg) = Rescued Food Weight (kg) × 2.5\n\nConverts every saved meal into live carbon footprint reduction displayed on the Impact Dashboard.")
    ]
    for i, (ititle, idesc) in enumerate(impacts):
        card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 5.95), Inches(1.8), Inches(5.75), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = DARK_CARD
        card.line.color.rgb = ACCENT_GREEN
        
        tf = card.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = ititle
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = GOLD
        
        p2 = tf.add_paragraph()
        p2.text = f"\n{idesc}"
        p2.font.size = Pt(13)
        p2.font.color.rgb = WHITE

    # -------------------------------------------------------------
    # SLIDE 8: Future Scope
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    apply_bg(slide8, BURGUNDY)
    add_header(slide8, 8, "Scaling PlateRelay Pan-India", "🔮 FUTURE SCOPE & ROADMAP")
    
    futures = [
        ("🏆 Volunteer Karma Points & Rewards", "• Concept: Volunteer runners earn PlateRelay Karma Points for every successful food delivery.\n\n• Why it helps: Volunteers redeem points for free fuel vouchers, public transport passes, or restaurant discounts."),
        ("🛵 Delivery Partner Fleet Integration", "• Concept: Utilizing existing delivery drivers (Swiggy/Zomato/Dunzo) during their slow off-peak hours (3 PM – 6 PM).\n\n• Why it helps: Provides zero-cost, reliable transportation for bulk surplus meals to shelters."),
        ("🤖 AI Surplus Demand Predictor", "• Concept: Machine Learning models analyzing past restaurant sales, weather & local event trends.\n\n• Why it helps: Predicts surplus food before cooking starts, pre-alerting NGOs in advance.")
    ]
    for i, (ftitle, fdesc) in enumerate(futures):
        card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 3.9), Inches(1.8), Inches(3.7), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = DARK_CARD
        card.line.color.rgb = GOLD
        
        tf = card.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = ftitle
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = GOLD
        
        p2 = tf.add_paragraph()
        p2.text = f"\n{fdesc}"
        p2.font.size = Pt(12)
        p2.font.color.rgb = WHITE

    # Save Presentation Output
    output_path1 = "c:/Users/sharm/OneDrive/Desktop/Code Blooded/PlateRelay_Hackathon_Presentation.pptx"
    output_path2 = "c:/Users/sharm/OneDrive/Desktop/Code Blooded/client/public/PlateRelay_Hackathon_Presentation.pptx"
    
    prs.save(output_path1)
    prs.save(output_path2)
    print(f"PowerPoint Presentation saved successfully to:\n1. {output_path1}\n2. {output_path2}")

if __name__ == "__main__":
    build_presentation()
